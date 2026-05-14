from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
import os

IMG = "images"

W = Inches(13.33)
H = Inches(7.5)

BLUE_DARK  = RGBColor(0x0F, 0x34, 0x60)
BLUE_MID   = RGBColor(0x31, 0x82, 0xCE)
BLUE_LIGHT = RGBColor(0x90, 0xCD, 0xF4)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY       = RGBColor(0x71, 0x80, 0x96)
DARK_TEXT  = RGBColor(0x1A, 0x36, 0x5D)
BODY_TEXT  = RGBColor(0x2D, 0x37, 0x48)
BG_ACCENT  = RGBColor(0xE8, 0xF4, 0xF8)

prs = Presentation()
prs.slide_width  = W
prs.slide_height = H

blank = prs.slide_layouts[6]  # completely blank

def add_slide():
    return prs.slides.add_slide(blank)

def rect(slide, l, t, w, h, fill=None, line=None):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.line.fill.background()
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def txbox(slide, text, l, t, w, h,
          size=18, bold=False, color=BODY_TEXT,
          align=PP_ALIGN.LEFT, wrap=True):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb

def heading(slide, text, top=Inches(0.4)):
    rect(slide, Inches(0.5), top + Inches(0.42), Inches(12.33), Pt(2), fill=BLUE_MID)
    txbox(slide, text,
          Inches(0.5), top, Inches(12), Inches(0.55),
          size=26, bold=True, color=DARK_TEXT)

def brand_bar(slide):
    rect(slide, 0, 0, W, Pt(6), fill=BLUE_MID)

def bg_white(slide):
    rect(slide, 0, 0, W, H, fill=WHITE)

def bg_accent(slide):
    rect(slide, 0, 0, W, H, fill=BG_ACCENT)

def bg_dark(slide):
    rect(slide, 0, 0, W, H, fill=BLUE_DARK)

def add_image(slide, path, l, t, w, h=None):
    if not os.path.exists(path):
        return
    if h:
        slide.shapes.add_picture(path, l, t, w, h)
    else:
        slide.shapes.add_picture(path, l, t, w)

def card(slide, l, t, w, h, title, body, title_color=BLUE_MID, bg=WHITE):
    rect(slide, l, t, w, h, fill=bg, line=RGBColor(0xE2, 0xE8, 0xF0))
    rect(slide, l, t, Pt(4), h, fill=BLUE_MID)
    txbox(slide, title, l+Pt(8), t+Inches(0.1), w-Pt(12), Inches(0.3),
          size=13, bold=True, color=title_color)
    txbox(slide, body, l+Pt(8), t+Inches(0.38), w-Pt(12), h-Inches(0.45),
          size=11, color=BODY_TEXT)

def card_dark(slide, l, t, w, h, title, body):
    rect(slide, l, t, w, h, fill=RGBColor(0x1A, 0x2A, 0x4A))
    rect(slide, l, t, Pt(4), h, fill=BLUE_LIGHT)
    txbox(slide, title, l+Pt(8), t+Inches(0.1), w-Pt(12), Inches(0.3),
          size=13, bold=True, color=BLUE_LIGHT)
    txbox(slide, body, l+Pt(8), t+Inches(0.38), w-Pt(12), h-Inches(0.45),
          size=11, color=RGBColor(0xE2, 0xE8, 0xF0))

# ─────────────────────────────────────────────
# SLIDE 1: タイトル
# ─────────────────────────────────────────────
s = add_slide()
rect(s, 0, 0, W, H, fill=BLUE_DARK)

txbox(s, "東大阪サミット_5月15日",
      Inches(1), Inches(0.5), Inches(11), Inches(0.5),
      size=16, bold=True, color=BLUE_LIGHT, align=PP_ALIGN.CENTER)

txbox(s, "カーフィルム",
      Inches(1), Inches(1.3), Inches(11), Inches(1.0),
      size=52, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

txbox(s, "Honest【オネスト】",
      Inches(1), Inches(2.3), Inches(11), Inches(0.8),
      size=40, bold=True, color=BLUE_LIGHT, align=PP_ALIGN.CENTER)

txbox(s, "あなたの愛車を、もっと快適に。",
      Inches(1), Inches(3.2), Inches(11), Inches(0.45),
      size=18, color=RGBColor(0xE2, 0xE8, 0xF0), align=PP_ALIGN.CENTER)

txbox(s, "酒井 修二",
      Inches(1), Inches(3.75), Inches(11), Inches(0.45),
      size=22, color=WHITE, align=PP_ALIGN.CENTER)

txbox(s, "東大阪市 ｜ 自動車カーフィルム専門店",
      Inches(1), Inches(4.25), Inches(11), Inches(0.35),
      size=13, color=GRAY, align=PP_ALIGN.CENTER)

add_image(s, f"{IMG}/image2.png", Inches(5.4), Inches(4.8), Inches(2.5))

# ─────────────────────────────────────────────
# SLIDE 2: 自己紹介
# ─────────────────────────────────────────────
s = add_slide()
bg_white(s)
brand_bar(s)
heading(s, "自己紹介")

txbox(s, "酒井 修二",
      Inches(0.6), Inches(1.15), Inches(5), Inches(0.45),
      size=22, bold=True, color=DARK_TEXT)
txbox(s, "1971年9月生まれ ｜ 東大阪市在住",
      Inches(0.6), Inches(1.6), Inches(5), Inches(0.3),
      size=12, color=GRAY)

timeline = [
    ("1992", "ガソリンスタンドで自動車業界へ"),
    ("1996", "俱楽部WALK 開業（従業員として）"),
    ("2001", "オネストとして独立開業"),
    ("2011", "守口門真青年会議所 卒業"),
    ("2018", "BNI シャングリラチャプター 入会"),
    ("2025", "東大阪サミット 参加"),
]
ty = Inches(2.05)
for year, event in timeline:
    txbox(s, year, Inches(0.6), ty, Inches(0.8), Inches(0.32),
          size=12, bold=True, color=BLUE_MID)
    rect(s, Inches(1.5), ty + Inches(0.12), Inches(0.12), Inches(0.12),
         fill=BLUE_MID)
    bold = year in ("2001", "2025")
    txbox(s, event, Inches(1.7), ty, Inches(4.1), Inches(0.32),
          size=13, bold=bold, color=BODY_TEXT)
    ty += Inches(0.42)

# モットーボックス
rect(s, Inches(6.8), Inches(1.15), Inches(6.0), Inches(5.5),
     fill=RGBColor(0xEB, 0xF8, 0xFF))
txbox(s, "モットー",
      Inches(7.0), Inches(1.3), Inches(5.6), Inches(0.35),
      size=13, bold=True, color=BLUE_MID)
txbox(s, "「決意した事を実行しないと\n心に負担が残っていく」\n\nお客様の「困った」を\n一緒に解決することを大切に、\n東大阪で30年以上活動しています。",
      Inches(7.0), Inches(1.7), Inches(5.6), Inches(3.5),
      size=14, color=RGBColor(0x2A, 0x4A, 0x6B))

# ─────────────────────────────────────────────
# SLIDE 3: Honest とは
# ─────────────────────────────────────────────
s = add_slide()
bg_white(s)
brand_bar(s)
heading(s, "Honest【オネスト】とは")

txbox(s, "東大阪を拠点に、カーフィルム・コーティングを専門とする自動車サービス店です。",
      Inches(0.6), Inches(1.1), Inches(12), Inches(0.35),
      size=13, color=RGBColor(0x4A, 0x55, 0x68))

cards3 = [
    ("🎯 専門特化",
     "カーフィルムとコーティングに特化。\n職人として一台一台丁寧に施工します。"),
    ("🤝 自動車業界の分業体制",
     "販売・整備・鈑金など各分野のプロが連携。\nオネストはフィルム・コーティングのスペシャリスト。"),
    ("📍 東大阪密着30年以上",
     "地域に根ざした信頼関係を大切に、\n口コミや紹介で多くのお客様にご利用いただいています。"),
    ("💡 幅広いネットワーク",
     "BNI・青年会議所など経営者ネットワークを活かし、\n最適な専門家をご紹介できます。"),
]
cx, cy = Inches(0.5), Inches(1.6)
cw, ch = Inches(6.1), Inches(2.3)
for i, (t, b) in enumerate(cards3):
    col = i % 2
    row = i // 2
    card(s, cx + col * Inches(6.5), cy + row * Inches(2.5), cw, ch, t, b)

# ─────────────────────────────────────────────
# SLIDE 4: 問題提起
# ─────────────────────────────────────────────
s = add_slide()
bg_accent(s)
brand_bar(s)

txbox(s, "車に乗っていて\n「暑いなぁ〜〜💦」\nと感じることはありませんか？",
      Inches(0.5), Inches(1.2), Inches(7.5), Inches(3.8),
      size=34, bold=True, color=DARK_TEXT, align=PP_ALIGN.CENTER)

txbox(s, "顔や腕に当たるジリジリとした日差し…\nその悩み、カーフィルムで解決できます。",
      Inches(0.5), Inches(5.2), Inches(7.5), Inches(0.8),
      size=14, color=GRAY, align=PP_ALIGN.CENTER)

add_image(s, f"{IMG}/image4.jpeg", Inches(8.3), Inches(0.8), Inches(4.5))
add_image(s, f"{IMG}/image5.jpeg", Inches(8.3), Inches(3.9), Inches(4.5))

# ─────────────────────────────────────────────
# SLIDE 5: カーフィルムでできること
# ─────────────────────────────────────────────
s = add_slide()
bg_white(s)
brand_bar(s)
heading(s, "カーフィルムでできること")

features = [
    "🌞  日射・遮熱効果で車内温度を下げる",
    "🕶️  外からの視界を遮りプライバシーを守る",
    "🛡️  紫外線（UV）から人・内装を保護する",
    "💥  ガラス破損時の飛散を防止する",
    "✨  ドレスアップ・見た目をスタイリッシュに",
]
fy = Inches(1.3)
for f in features:
    rect(s, Inches(0.55), fy + Inches(0.06), Inches(0.22), Inches(0.22),
         fill=BLUE_MID)
    txbox(s, f, Inches(0.9), fy, Inches(5.8), Inches(0.36),
          size=15, color=BODY_TEXT)
    fy += Inches(0.48)

rect(s, Inches(0.55), fy + Inches(0.1), Inches(6.1), Inches(1.1),
     fill=RGBColor(0xEB, 0xF8, 0xFF))
txbox(s, "施工の流れ",
      Inches(0.7), fy + Inches(0.15), Inches(5.9), Inches(0.3),
      size=12, bold=True, color=BLUE_MID)
txbox(s, "① ヒアリング・目的確認　② フィルム選定・ご提案　③ 丁寧な施工　④ 内装保護・仕上げ確認",
      Inches(0.7), fy + Inches(0.45), Inches(5.9), Inches(0.5),
      size=11, color=RGBColor(0x2A, 0x4A, 0x6B))

add_image(s, f"{IMG}/image8.jpeg", Inches(7.1), Inches(1.1), Inches(5.8))

# ─────────────────────────────────────────────
# SLIDE 6: カーコーティングの魅力
# ─────────────────────────────────────────────
s = add_slide()
bg_white(s)
brand_bar(s)
heading(s, "カーコーティングの魅力")

txbox(s, "「いつも綺麗な車に乗っていたい♬」そのご要望にお応えします。",
      Inches(0.6), Inches(1.1), Inches(12), Inches(0.35),
      size=13, color=RGBColor(0x4A, 0x55, 0x68))

coating = [
    ("🚿  洗車が楽になる",   "汚れがつきにくく、水洗いだけで綺麗に"),
    ("💎  輝きが長続き",     "深みのある光沢で新車のような美しさを維持"),
    ("🛡️  ボディを保護",    "紫外線・酸性雨・飛び石から塗装を守る"),
]
for i, (t, b) in enumerate(coating):
    cx2 = Inches(0.5) + i * Inches(4.27)
    rect(s, cx2, Inches(1.55), Inches(4.0), Inches(2.8),
         fill=WHITE, line=RGBColor(0xE2, 0xE8, 0xF0))
    rect(s, cx2, Inches(1.55), Inches(4.0), Pt(4), fill=BLUE_MID)
    txbox(s, t, cx2 + Inches(0.1), Inches(1.75), Inches(3.8), Inches(0.4),
          size=14, bold=True, color=BLUE_MID, align=PP_ALIGN.CENTER)
    txbox(s, b, cx2 + Inches(0.1), Inches(2.25), Inches(3.8), Inches(0.8),
          size=12, color=BODY_TEXT, align=PP_ALIGN.CENTER)

rect(s, Inches(0.5), Inches(4.55), Inches(12.33), Inches(0.9),
     fill=RGBColor(0xEB, 0xF8, 0xFF))
txbox(s, "コーティングはお車の「投資」です。査定額アップにもつながり、長期的にお得なカーケアです。",
      Inches(0.7), Inches(4.65), Inches(11.9), Inches(0.7),
      size=14, bold=False, color=RGBColor(0x2A, 0x4A, 0x6B), align=PP_ALIGN.CENTER)

add_image(s, f"{IMG}/image15.jpg", Inches(0.5), Inches(5.6), Inches(12.33), Inches(1.55))

# ─────────────────────────────────────────────
# SLIDE 7: サービス一覧
# ─────────────────────────────────────────────
s = add_slide()
bg_accent(s)
brand_bar(s)
heading(s, "サービス一覧")

services = [
    ("🎞️", "カーフィルム施工"),
    ("✨", "カーコーティング"),
    ("🔍", "車検・一般整備"),
    ("🛒", "車の販売・買取"),
    ("🔧", "パーツ交換作業"),
    ("🎨", "鈑金・塗装"),
]
for i, (icon, name) in enumerate(services):
    col = i % 2
    row = i // 2
    sx = Inches(0.6) + col * Inches(6.5)
    sy = Inches(1.55) + row * Inches(1.5)
    rect(s, sx, sy, Inches(6.1), Inches(1.2),
         fill=WHITE, line=RGBColor(0xE2, 0xE8, 0xF0))
    txbox(s, icon, sx + Inches(0.2), sy + Inches(0.25), Inches(0.6), Inches(0.6),
          size=24, align=PP_ALIGN.CENTER)
    txbox(s, name, sx + Inches(0.9), sy + Inches(0.3), Inches(4.9), Inches(0.55),
          size=18, bold=True, color=BODY_TEXT)

txbox(s, "※ 自動車業界の分業体制を活かし、専門パートナーと連携してワンストップで対応します",
      Inches(0.6), Inches(6.8), Inches(12), Inches(0.35),
      size=11, color=GRAY, align=PP_ALIGN.CENTER)

# ─────────────────────────────────────────────
# SLIDE 8: 過去の実績
# ─────────────────────────────────────────────
s = add_slide()
bg_white(s)
brand_bar(s)
heading(s, "過去の実績・強み")

achievements = [
    ("24年", "独立開業からの実績年数\n（2001年〜）"),
    ("30年以上", "自動車業界での\nキャリア年数"),
    ("地域密着", "東大阪を拠点に\n口コミ・紹介で拡大"),
]
for i, (num, label) in enumerate(achievements):
    ax = Inches(0.5) + i * Inches(4.27)
    rect(s, ax, Inches(1.2), Inches(4.0), Inches(1.7),
         fill=WHITE, line=RGBColor(0xE2, 0xE8, 0xF0))
    txbox(s, num, ax, Inches(1.3), Inches(4.0), Inches(0.7),
          size=30, bold=True, color=BLUE_MID, align=PP_ALIGN.CENTER)
    txbox(s, label, ax, Inches(2.0), Inches(4.0), Inches(0.7),
          size=11, color=GRAY, align=PP_ALIGN.CENTER)

rect(s, Inches(0.5), Inches(3.1), Inches(12.33), Inches(1.4),
     fill=RGBColor(0xEB, 0xF8, 0xFF))
txbox(s, "お客様の声",
      Inches(0.7), Inches(3.2), Inches(12), Inches(0.3),
      size=12, bold=True, color=BLUE_MID)
txbox(s, "「フィルムを貼ってから夏の車内が全然違う！快適になりました」\n「丁寧な施工で仕上がりに大満足。友人にも紹介しました」\n「相談しやすく、予算に合わせて提案してくれるので安心して任せられます」",
      Inches(0.7), Inches(3.5), Inches(12), Inches(0.9),
      size=11, color=RGBColor(0x2A, 0x4A, 0x6B))

photos = ["image9.jpeg", "image11.jpeg", "image12.jpg", "image13.jpg"]
pw = Inches(3.0)
for i, p in enumerate(photos):
    add_image(s, f"{IMG}/{p}", Inches(0.5) + i * Inches(3.2), Inches(4.65), pw, Inches(2.5))

# ─────────────────────────────────────────────
# SLIDE 9: 社会への取り組み
# ─────────────────────────────────────────────
s = add_slide()
bg_dark(s)
heading_text = "社会への取り組み"
txbox(s, heading_text,
      Inches(0.5), Inches(0.4), Inches(12), Inches(0.55),
      size=26, bold=True, color=WHITE)
rect(s, Inches(0.5), Inches(0.88), Inches(12.33), Pt(2), fill=BLUE_LIGHT)

card_dark(s, Inches(0.5), Inches(1.05), Inches(12.33), Inches(5.7),
          "🎗️ 色素性乾皮症（指定難病159）の方々へのサポート",
          "色素性乾皮症は紫外線の影響を強く受け、日常生活に大きな支障をきたす難病です。\n\n"
          "カーフィルムによる紫外線カットは、こうした方々の外出負担を大幅に軽減できます。\n"
          "費用面での負担を軽減する方法もご案内できますので、お困りの方はご相談ください。")

# ─────────────────────────────────────────────
# SLIDE 10: まとめ・クロージング
# ─────────────────────────────────────────────
s = add_slide()
bg_dark(s)

txbox(s, "快適なカーライフを\nお手伝いします♪",
      Inches(1), Inches(0.8), Inches(11), Inches(1.8),
      size=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

txbox(s, "車のことで困ったら、まずご相談ください。",
      Inches(1), Inches(2.7), Inches(11), Inches(0.5),
      size=18, color=BLUE_LIGHT, align=PP_ALIGN.CENTER)

rect(s, Inches(3.5), Inches(3.35), Inches(6.33), Inches(2.2),
     fill=RGBColor(0x1A, 0x2A, 0x4A))
txbox(s, "Honest【オネスト】",
      Inches(3.7), Inches(3.45), Inches(5.9), Inches(0.4),
      size=16, bold=True, color=BLUE_LIGHT, align=PP_ALIGN.CENTER)
txbox(s, "代表　酒井 修二\n📍 東大阪市\n🚗 カーフィルム専門店",
      Inches(3.7), Inches(3.9), Inches(5.9), Inches(1.4),
      size=14, color=RGBColor(0xE2, 0xE8, 0xF0), align=PP_ALIGN.CENTER)

add_image(s, f"{IMG}/image18.png", Inches(6.1), Inches(5.75), Inches(1.1))

txbox(s, "ご清聴ありがとうございました。",
      Inches(1), Inches(6.9), Inches(11), Inches(0.35),
      size=13, color=GRAY, align=PP_ALIGN.CENTER)

# ─────────────────────────────────────────────
out = "honest_presentation.pptx"
prs.save(out)
print(f"✅ 保存完了: {out}")
